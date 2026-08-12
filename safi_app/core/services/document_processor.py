"""
Document text extraction service.

Extracts plain text from uploaded files (PDF, DOCX, XLSX, TXT, MD, CSV, and
images via OCR) so it can be injected as context into the user's prompt.

IMAGES ARE OCR'd TO TEXT — the image itself is never sent to a model. That is a
governance decision, not a limitation of convenience: Phase Zero is a literal
substring scan over the prompt and cannot see inside an image, so attaching image
bytes to a model call would create an input channel the deterministic tier
structurally cannot inspect. Turning the image into text first means it travels
the same gated path as every other attachment. See GOVERNANCE_BACKLOG 32z for the
sketch of how real vision could be governed, and why it is not built.
"""
import os
import csv
import io
import logging
from typing import Tuple

log = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.tiff', '.tif', '.webp', '.bmp'}
ALLOWED_EXTENSIONS = {'.txt', '.md', '.pdf', '.docx', '.xlsx', '.csv'} | IMAGE_EXTENSIONS

# A page that yields fewer than this many characters is treated as scanned rather
# than as genuinely sparse, and re-read with OCR when OCR is available. Chosen to
# clear a stray header or page number without swallowing a real, short page.
_SCANNED_PAGE_CHAR_FLOOR = 40


def allowed_file(filename: str) -> bool:
    """Checks if the file extension is in the allowed set."""
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS


def extract_text(file_storage, filename: str, max_chars: int = 50000) -> Tuple[str, int]:
    """
    Extracts text from an uploaded file.

    Args:
        file_storage: A file-like object (e.g., from Flask's request.files).
        filename: The original filename (used to detect format).
        max_chars: Maximum characters to return. Documents exceeding this
                   will be truncated with a notice.

    Returns:
        A tuple of (extracted_text, total_chars_before_truncation).
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext in ('.txt', '.md'):
        text = file_storage.read().decode('utf-8', errors='replace')
    elif ext == '.csv':
        text = _extract_csv(file_storage)
    elif ext == '.pdf':
        text = _extract_pdf(file_storage)
    elif ext == '.docx':
        text = _extract_docx(file_storage)
    elif ext == '.xlsx':
        text = _extract_xlsx(file_storage)
    elif ext in IMAGE_EXTENSIONS:
        text = _extract_image(file_storage, filename)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    total_chars = len(text)
    if total_chars > max_chars:
        text = text[:max_chars] + (
            f"\n\n[... Document truncated at {max_chars:,} characters. "
            f"Total: {total_chars:,} characters ...]"
        )

    return text, total_chars


def _extract_csv(file_storage) -> str:
    """Reads a CSV and formats it as a Markdown table."""
    content = file_storage.read().decode('utf-8', errors='replace')
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    if not rows:
        return ""

    header = rows[0]
    lines = ["| " + " | ".join(header) + " |"]
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in rows[1:]:
        # Pad or truncate columns to match header length
        padded = row + [''] * (len(header) - len(row))
        lines.append("| " + " | ".join(padded[:len(header)]) + " |")
    return "\n".join(lines)


def _extract_pdf(file_storage) -> str:
    """Extracts text from a PDF using PyPDF2."""
    try:
        import PyPDF2
    except ImportError:
        raise ValueError(
            "PDF support requires PyPDF2. "
            "Install with: pip install PyPDF2"
        )

    reader = PyPDF2.PdfReader(file_storage)
    pages = []
    ocr_pages = 0
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()

        # A scanned page has an embedded image and no text layer, so PyPDF2
        # returns nothing and the page silently disappears from the extraction.
        # Before OCR existed, a fully scanned PDF raised "may be image-based" and
        # a PARTLY scanned one was worse: it succeeded, returned only the typed
        # pages, and looked complete.
        if len(text) < _SCANNED_PAGE_CHAR_FLOOR:
            ocr_text = _ocr_pdf_page_images(page)
            if ocr_text:
                text = ocr_text
                ocr_pages += 1

        if text:
            pages.append(f"--- Page {i + 1} ---\n{text}")

    if not pages:
        raise ValueError(
            "Could not extract any text from this PDF. If it is a scan, OCR is "
            "unavailable or found no readable text."
        )

    if ocr_pages:
        log.info("PDF extraction used OCR for %d of %d page(s).", ocr_pages, len(reader.pages))
        pages.append(
            f"[NOTE: {ocr_pages} of {len(reader.pages)} pages in this document had no text "
            "layer and were read by OCR. OCR output can contain recognition errors; treat "
            "exact figures and names from those pages as uncertain and say so if they matter.]"
        )

    return "\n\n".join(pages)


def _ocr_available() -> bool:
    """True when both the Python binding and the tesseract binary are present.
    The wheel installs without the binary, so importing pytesseract proves
    nothing on its own."""
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image_bytes(data: bytes) -> str:
    """OCR raw image bytes. Returns '' on any failure — a caller reading a
    scanned page must be able to fall through to whatever it had."""
    try:
        import pytesseract
        from PIL import Image
        with Image.open(io.BytesIO(data)) as img:
            # Some scans arrive as 1-bit or palette images, which tesseract reads
            # poorly; RGB is the reliable input.
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            return (pytesseract.image_to_string(img) or "").strip()
    except Exception as e:
        log.warning("OCR failed on an image (%s: %s)", type(e).__name__, e)
        return ""


def _ocr_pdf_page_images(page) -> str:
    """
    OCR the images embedded in one PDF page.

    Deliberately reads the page's own XObject images rather than rasterising the
    page: rasterising needs Poppler or PyMuPDF, neither of which is a dependency,
    and a scanned page is in practice one full-page image. A page whose text is
    drawn as vectors will not be recovered this way — it is a limitation, not a
    bug, and the extraction note tells the model OCR was involved.
    """
    if not _ocr_available():
        return ""
    out = []
    try:
        images = list(getattr(page, "images", []) or [])
    except Exception:
        return ""
    for img in images:
        try:
            text = _ocr_image_bytes(img.data)
        except Exception:
            continue
        if text:
            out.append(text)
    return "\n".join(out).strip()


def _extract_image(file_storage, filename: str) -> str:
    """Extracts text from an image via OCR."""
    if not _ocr_available():
        raise ValueError(
            "Image support requires OCR, which is not available in this deployment. "
            "Install the tesseract-ocr system package and the pytesseract module."
        )
    text = _ocr_image_bytes(file_storage.read())
    if not text:
        raise ValueError(
            f"No readable text was found in {os.path.basename(filename)}. "
            "Images are read by OCR, so photographs, diagrams and charts without "
            "text will come back empty."
        )
    # The provenance matters to whoever reads the answer: OCR output is a
    # transcription, not the document, and the model should not present a
    # misrecognised figure as if it were quoted.
    return (
        f"[OCR TRANSCRIPTION of the image {os.path.basename(filename)}. This text was "
        "recognised from pixels and may contain errors; treat exact figures, names and "
        "codes as uncertain, and say so if the answer depends on them.]\n\n" + text
    )


def _extract_docx(file_storage) -> str:
    """Extracts text from a DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ValueError(
            "DOCX support requires python-docx. "
            "Install with: pip install python-docx"
        )

    doc = Document(file_storage)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    if not paragraphs:
        raise ValueError("Could not extract any text from this DOCX file.")

    return "\n\n".join(paragraphs)


def _extract_xlsx(file_storage) -> str:
    """Extracts cell values from an XLSX workbook as one Markdown table per sheet."""
    try:
        import openpyxl
    except ImportError:
        raise ValueError(
            "XLSX support requires openpyxl. "
            "Install with: pip install openpyxl"
        )

    wb = openpyxl.load_workbook(file_storage, read_only=True, data_only=True)
    sheets = []

    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            # Skip fully empty rows; stringify cells, blanking out None.
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                rows.append(cells)

        if not rows:
            continue

        width = max(len(r) for r in rows)
        header = rows[0] + [''] * (width - len(rows[0]))
        lines = ["| " + " | ".join(header) + " |"]
        lines.append("| " + " | ".join(["---"] * width) + " |")
        for row in rows[1:]:
            padded = row + [''] * (width - len(row))
            lines.append("| " + " | ".join(padded) + " |")

        sheets.append(f"### Sheet: {ws.title}\n\n" + "\n".join(lines))

    wb.close()

    if not sheets:
        raise ValueError("Could not extract any data from this XLSX file.")

    return "\n\n".join(sheets)
